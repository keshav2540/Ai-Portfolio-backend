"""
Test Suite for AI Portfolio Backend (Groq LLM + Embeddings + FastAPI Integration)
Run with: python test_suite.py
"""
import io
import json
import logging
import os
import sys
import unittest
from unittest.mock import MagicMock, patch

os.environ["TESTING"] = "1"
os.environ["GROQ_API_KEY"] = "gsk_test_mock_groq_key_12345"
os.environ["GROQ_MODEL"] = "llama-3.3-70b-versatile"
os.environ["GROQ_FALLBACK_MODEL"] = "llama-3.1-8b-instant"
os.environ["GEMINI_API_KEY"] = "test-secret-gemini-key-12345"

sys.path.insert(0, os.path.dirname(__file__))

# Suppress logs during tests
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("httpcore").setLevel(logging.WARNING)
logging.getLogger("langchain").setLevel(logging.WARNING)
logging.getLogger("langchain_groq").setLevel(logging.WARNING)
logging.getLogger("groq").setLevel(logging.WARNING)

import store
store.messages = None
store.matches = None
store.init = lambda: None

import llm
import main
import embeddings
import matcher


class TestGroqLLM(unittest.TestCase):
    def setUp(self):
        llm._response_cache.clear()
        llm.RETRY_BACKOFF_DELAYS = [0.001, 0.001]

    def tearDown(self):
        llm._response_cache.clear()

    def test_groq_ask_success(self):
        """ask() invokes primary Groq model and returns trimmed text."""
        mock_client = MagicMock()
        mock_client.invoke.return_value = MagicMock(content="  Portfolio response from Groq  ")

        with patch("llm._create_groq_client", return_value=mock_client):
            res = llm.ask("system prompt", [{"role": "user", "content": "hello"}])
            self.assertEqual(res, "Portfolio response from Groq")
            mock_client.invoke.assert_called_once()

    def test_groq_primary_fails_falls_back_to_secondary_model(self):
        """Primary model fails retries -> falls back to secondary Groq model and succeeds."""
        clients_created = []

        def mock_client_factory(model_name, **kwargs):
            client = MagicMock()
            clients_created.append(model_name)
            if model_name == "llama-3.3-70b-versatile":
                client.invoke.side_effect = RuntimeError("Primary 503 Service Unavailable")
            else:
                client.invoke.return_value = MagicMock(content="Fallback model response")
            return client

        with patch("llm._create_groq_client", side_effect=mock_client_factory):
            res = llm.ask("system prompt", [{"role": "user", "content": "hello"}])
            self.assertEqual(res, "Fallback model response")
            # Primary model tried twice (attempts 1 & 2), then fallback tried
            self.assertIn("llama-3.3-70b-versatile", clients_created)
            self.assertIn("llama-3.1-8b-instant", clients_created)

    def test_groq_all_models_fail_raises_clear_error(self):
        """When all models and retries fail, raises RuntimeError including underlying error."""
        def mock_client_factory(model_name, **kwargs):
            client = MagicMock()
            client.invoke.side_effect = RuntimeError("RateLimitError 429")
            return client

        with patch("llm._create_groq_client", side_effect=mock_client_factory):
            with self.assertRaises(RuntimeError) as ctx:
                llm.ask("system prompt", [{"role": "user", "content": "hello"}])
            self.assertIn("Groq service unavailable across all models", str(ctx.exception))
            self.assertIn("RateLimitError 429", str(ctx.exception))

    def test_groq_stream_success(self):
        """ask_stream() yields plain text chunks."""
        from langchain_core.messages import AIMessageChunk

        def fake_stream(messages, **kwargs):
            for part in ["Hello", " ", "world", "!"]:
                yield AIMessageChunk(content=part)

        mock_client = MagicMock()
        mock_client.stream.side_effect = fake_stream

        with patch("llm._create_groq_client", return_value=mock_client):
            chunks = list(llm.ask_stream("system prompt", [{"role": "user", "content": "hello"}]))
            self.assertEqual(chunks, ["Hello", " ", "world", "!"])

    def test_groq_stream_mid_stream_error_does_not_retry(self):
        """Tokens already sent to user mid-stream must NOT be retried (avoids duplicate output)."""
        from langchain_core.messages import AIMessageChunk

        def fake_stream_interrupted(messages, **kwargs):
            yield AIMessageChunk(content="First chunk")
            raise RuntimeError("Stream connection dropped")

        mock_client = MagicMock()
        mock_client.stream.side_effect = fake_stream_interrupted

        with patch("llm._create_groq_client", return_value=mock_client):
            gen = llm.ask_stream("system prompt", [{"role": "user", "content": "hello"}])
            first = next(gen)
            self.assertEqual(first, "First chunk")
            with self.assertRaises(RuntimeError) as ctx:
                next(gen)
            self.assertIn("interrupted", str(ctx.exception).lower())

    def test_groq_response_caching(self):
        """Identical non-json queries are served from 10-minute cache."""
        mock_client = MagicMock()
        mock_client.invoke.return_value = MagicMock(content="Cached response text")

        with patch("llm._create_groq_client", return_value=mock_client):
            r1 = llm.ask("system", [{"role": "user", "content": "same query"}])
            r2 = llm.ask("system", [{"role": "user", "content": "same query"}])
            self.assertEqual(r1, "Cached response text")
            self.assertEqual(r2, "Cached response text")
            self.assertEqual(mock_client.invoke.call_count, 1)

    def test_groq_ask_json_strips_markdown_fences(self):
        """ask_json() strips markdown code fences and handles trailing commas."""
        mock_client = MagicMock()
        mock_client.invoke.return_value = MagicMock(
            content='```json\n{\n  "score": 92,\n  "status": "fit",\n}\n```'
        )

        with patch("llm._create_groq_client", return_value=mock_client):
            data = llm.ask_json("system", "evaluate candidate")
            self.assertEqual(data.get("score"), 92)
            self.assertEqual(data.get("status"), "fit")


class TestEmbeddings(unittest.TestCase):
    def test_testing_mode_returns_768_dim_vector(self):
        """Embeddings module outputs 768-dim vector matching Qdrant collection size."""
        vec = embeddings.get_embedding("test text")
        self.assertEqual(len(vec), 768)
        self.assertEqual(embeddings.EMBEDDING_DIM, 768)


class TestMatcher(unittest.TestCase):
    """Task 5: Unit tests for matcher.py handling dict, list, wrapped dict, and invalid responses."""

    def test_matcher_dict_response(self):
        """Standard dict response from ask_json parses successfully."""
        mock_req = {"title": "React Engineer", "must_have": ["React", "CSS"], "nice_to_have": ["TypeScript"]}
        mock_cmp = {
            "must_have": [{"item": "React", "level": "yes", "evidence": "Used React"}],
            "nice_to_have": [],
            "experience": "yes",
            "reason": "Strong frontend background.",
        }

        def mock_ask_json(sys_prompt, prompt, max_tokens=None):
            if "job descriptions" in sys_prompt.lower():
                return mock_req
            return mock_cmp

        with patch("matcher.ask_json", side_effect=mock_ask_json):
            res = matcher.run("Keshav has 4 years experience in React and CSS", "Senior React Developer")
            self.assertEqual(res["title"], "React Engineer")
            self.assertIn("score", res)
            self.assertIn(res["verdict"], ["fit", "partial", "not_fit"])
            self.assertNotIn("note", res)

    def test_matcher_list_response(self):
        """LLM returns a list of job descriptions; matcher extracts first JD and marks multiple JDs note."""
        mock_list = [
            {"title": "Role 1: Frontend", "must_have": ["React"], "nice_to_have": []},
            {"title": "Role 2: Backend", "must_have": ["Python"], "nice_to_have": []},
        ]
        mock_cmp = {
            "must_have": [{"item": "React", "level": "yes", "evidence": "React exp"}],
            "nice_to_have": [],
            "experience": "yes",
            "reason": "Good match.",
        }

        def mock_ask_json(sys_prompt, prompt, max_tokens=None):
            if "job descriptions" in sys_prompt.lower():
                return mock_list
            return mock_cmp

        with patch("matcher.ask_json", side_effect=mock_ask_json):
            res = matcher.run("Resume text with React", "Job 1: Frontend. Job 2: Backend.")
            self.assertEqual(res["title"], "Role 1: Frontend")
            self.assertIn("score", res)
            self.assertEqual(res.get("note"), "Multiple job descriptions found, matched against the first one.")

    def test_matcher_wrapped_dict_response(self):
        """LLM returns a wrapped dict e.g. {'job_descriptions': [...]}; matcher extracts first JD."""
        mock_wrapped = {
            "job_descriptions": [
                {"title": "Cloud Architect", "must_have": ["AWS", "Terraform"], "nice_to_have": []},
                {"title": "DevOps Engineer", "must_have": ["Kubernetes"], "nice_to_have": []},
            ]
        }
        mock_cmp = {
            "must_have": [{"item": "AWS", "level": "yes", "evidence": "Certified AWS"}],
            "nice_to_have": [],
            "experience": "yes",
            "reason": "Cloud experience matches.",
        }

        def mock_ask_json(sys_prompt, prompt, max_tokens=None):
            if "job descriptions" in sys_prompt.lower():
                return mock_wrapped
            return mock_cmp

        with patch("matcher.ask_json", side_effect=mock_ask_json):
            res = matcher.run("Resume text with AWS", "Job 1 ... Job 2 ...")
            self.assertEqual(res["title"], "Cloud Architect")
            self.assertIn("score", res)
            self.assertEqual(res.get("note"), "Multiple job descriptions found, matched against the first one.")

    def test_matcher_empty_or_invalid_response_raises_value_error(self):
        """Invalid or empty response from ask_json raises ValueError."""
        for invalid_val in [None, "", [], [1, 2, 3], {"unexpected": "no_must_have"}]:
            with patch("matcher.ask_json", return_value=invalid_val):
                with self.assertRaises(ValueError) as ctx:
                    matcher.run("Resume text", "Job Description Text")
                self.assertIn("Could not", str(ctx.exception))


class TestFastAPIIntegration(unittest.TestCase):
    def test_root_endpoint_returns_online(self):
        res = main.root()
        self.assertEqual(res["status"], "online")
        self.assertIn("health", res)

    def test_match_returns_503_on_upstream_failure(self):
        dummy_req = MagicMock()
        dummy_req.client.host = "127.0.0.1"
        with patch("main.rag.get_all_indexed_resume", return_value="Resume content"), \
             patch("main.matcher.run", side_effect=RuntimeError("Groq service unavailable across all models")):
            with self.assertRaises(main.HTTPException) as ctx:
                main.match(dummy_req, session_id="test-session-12345", file=None, jd_text="React dev role " * 10)
            self.assertEqual(ctx.exception.status_code, 503)
            self.assertIn("Abhi AI service busy hai", ctx.exception.detail)

    def test_chat_non_stream_returns_503_on_upstream_failure(self):
        dummy_req = MagicMock()
        dummy_req.client.host = "127.0.0.1"
        body = main.ChatIn(session_id="test-session-12345", message="hello")
        with patch("main.get_resume_context", return_value="Resume content"), \
             patch("main.ask", side_effect=RuntimeError("Groq service unavailable across all models")):
            with self.assertRaises(main.HTTPException) as ctx:
                main.chat(body, dummy_req)
            self.assertEqual(ctx.exception.status_code, 503)
            self.assertIn("Abhi AI service busy hai", ctx.exception.detail)

    def test_stream_generator_catches_error_and_yields_friendly_message(self):
        with patch("main.ask_stream", side_effect=RuntimeError("Groq streaming connection dropped")), \
             patch("main.get_resume_context", return_value="Mock resume"):
            body = main.ChatIn(session_id="test-session-12345", message="test")
            dummy_request = MagicMock()
            dummy_request.client.host = "127.0.0.1"

            resp = main.chat_stream(body, dummy_request)
            sync_gen = resp.body_iterator.ag_frame.f_locals.get("iterator", resp.body_iterator)
            chunks = list(sync_gen)
            full_text = "".join(chunks)
            self.assertIn("Abhi AI service busy hai", full_text)

    def test_chat_empty_or_none_message_handled_safely(self):
        """Empty or None message in /api/chat returns friendly prompt without crashing."""
        dummy_req = MagicMock()
        dummy_req.client.host = "127.0.0.1"
        res1 = main.chat(main.ChatIn(session_id="test-session-12345", message=""), dummy_req)
        self.assertIn("sawal", res1["reply"].lower())
        res2 = main.chat(main.ChatIn(session_id="test-session-12345", message=None), dummy_req)
        self.assertIn("sawal", res2["reply"].lower())

    def test_chat_stream_empty_or_none_message_handled_safely(self):
        """Empty or None message in /api/chat/stream streams prompt without crashing."""
        dummy_req = MagicMock()
        dummy_req.client.host = "127.0.0.1"
        resp = main.chat_stream(main.ChatIn(session_id="test-session-12345", message="  "), dummy_req)
        sync_gen = resp.body_iterator.ag_frame.f_locals.get("iterator", resp.body_iterator)
        chunks = list(sync_gen)
        self.assertIn("sawal", "".join(chunks).lower())


class TestJDParser(unittest.TestCase):
    def test_docx_extraction_paragraphs_and_tables(self):
        """DOCX with paragraphs and tables extracts text correctly."""
        import jd_parser
        from docx import Document
        import io

        doc = Document()
        doc.add_paragraph("Job Title: Software Engineer")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Skill"
        table.rows[0].cells[1].text = "Python"
        buf = io.BytesIO()
        doc.save(buf)

        text = jd_parser.extract("sample.docx", buf.getvalue())
        self.assertIn("Software Engineer", text)
        self.assertIn("Python", text)

    def test_unsupported_format_raises_value_error(self):
        """Unsupported file extensions raise clear ValueError."""
        import jd_parser
        with self.assertRaises(ValueError) as ctx:
            jd_parser.extract("test.exe", b"invalid data")
        self.assertIn("Unsupported file format", str(ctx.exception))

    def test_image_extraction_invokes_vision(self):
        """Images invoke Groq vision extraction."""
        import jd_parser
        with patch("jd_parser._extract_text_from_image", return_value="Senior Developer Needed\nRequirements: Python"):
            text = jd_parser.extract("job.png", b"fake_png_data")
            self.assertIn("Senior Developer Needed", text)



if __name__ == "__main__":
    unittest.main(verbosity=2)
