"""Turns an uploaded PDF/DOCX into clean text. The file is read in memory and never saved."""
import io, re
from pypdf import PdfReader
from docx import Document


def _clean(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", l).strip() for l in text.splitlines()]
    lines = [l for l in lines if l and not l.isdigit()]  # drop blanks and bare page numbers
    return "\n".join(lines)


def extract(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Only PDF and DOCX files are supported.")
    return _clean(text)
